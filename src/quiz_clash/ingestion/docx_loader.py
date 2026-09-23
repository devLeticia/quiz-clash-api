import io

from docx import Document as DocxDocument
from langchain_core.documents import Document

from quiz_clash.core.exceptions import DocxProcessingError, EmptyDocumentError


def load_docx(file_bytes: bytes, source_name: str = "uploaded_document") -> list[Document]:
    """Extract text from a DOCX file."""
    try:
        docx_file = DocxDocument(io.BytesIO(file_bytes))
        text = "\n".join(p.text for p in docx_file.paragraphs)
        print(f"o texto: {text}")
    except Exception as e:
        raise DocxProcessingError(f"Could not read the DOCX file: {e}") from e

    if not text.strip():
        raise EmptyDocumentError("No extractable text found in the DOCX file.")

    return [Document(page_content=text, metadata={"source": source_name})]
