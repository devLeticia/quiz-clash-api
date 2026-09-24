import io

from langchain_core.documents import Document
from pptx import Presentation

from quiz_clash.core.exceptions import EmptyDocumentError, PptxProcessingError


def load_pptx(file_bytes: bytes, source_name: str = "uploaded_presentation") -> list[Document]:
    """Extract text from a PPTX file (slide titles, text boxes, bullet points)."""
    try:
        presentation = Presentation(io.BytesIO(file_bytes))
        text_parts = []

        for slide in presentation.slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        line = "".join(run.text for run in paragraph.runs)
                        if line.strip():
                            text_parts.append(line)

        text = "\n".join(text_parts)
    except Exception as e:
        raise PptxProcessingError(f"Could not read the PPTX file: {e}") from e

    if not text.strip():
        raise EmptyDocumentError("No extractable text found in the PPTX file.")

    return [Document(page_content=text, metadata={"source": source_name})]
